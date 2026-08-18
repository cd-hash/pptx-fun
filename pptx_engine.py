from pptx import Presentation
from pptx.util import Pt

# --- Helper Methods ---

def get_slide_by_title(prs, target_title):
    for slide in prs.slides:
        if slide.shapes.title and slide.shapes.title.has_text_frame:
            if slide.shapes.title.text_frame.text.strip() == target_title:
                return slide
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip() == target_title:
                return slide
    return None

def write_to_cell_preserving_format(cell, value):
    p = cell.text_frame.paragraphs[0]
    if p.runs:
        p.runs[0].text = str(value)
        for extra_run in p.runs[1:]: extra_run.text = ""
    else:
        cell.text = str(value)

# --- Standard Text Replacement ---

def process_text_frame(text_frame, replace_dict):
    for paragraph in text_frame.paragraphs:
        for key, value in replace_dict.items():
            for run in paragraph.runs:
                if key in run.text: run.text = run.text.replace(key, str(value))
            
            while key in paragraph.text:
                runs = paragraph.runs
                run_texts = [r.text for r in runs]
                full_text = "".join(run_texts)
                start_idx = full_text.find(key)
                if start_idx == -1: break 
                end_idx = start_idx + len(key)
                
                def get_run_and_offset(abs_idx):
                    current = 0
                    for i, run_text in enumerate(run_texts):
                        if current <= abs_idx < current + len(run_text): return i, abs_idx - current
                        current += len(run_text)
                    return len(run_texts) - 1, len(run_texts[-1])
                
                start_run_idx, start_local = get_run_and_offset(start_idx)
                end_run_idx, end_local = get_run_and_offset(end_idx)
                
                runs[start_run_idx].text = run_texts[start_run_idx][:start_local] + str(value)
                for i in range(start_run_idx + 1, end_run_idx): runs[i].text = ""
                
                if start_run_idx != end_run_idx: runs[end_run_idx].text = run_texts[end_run_idx][end_local:]
                else: runs[start_run_idx].text += run_texts[end_run_idx][end_local:]

def process_shape(shape, replace_dict):
    if shape.has_text_frame: process_text_frame(shape.text_frame, replace_dict)
    elif shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells: process_text_frame(cell.text_frame, replace_dict)
    elif hasattr(shape, 'shape_type') and shape.shape_type == 6:
        for sub_shape in shape.shapes: process_shape(sub_shape, replace_dict)

# --- Table Injection Methods ---

def fill_table_from_df(slide, df, include_headers=True, start_row=0, start_col=0, table_name=None, table_index=0):
    table_shape = None
    current_table_idx = 0
    
    for shape in slide.shapes:
        if shape.has_table:
            # Route A: Target by specific name
            if table_name:
                if table_name in " ".join(cell.text_frame.text.strip() for cell in shape.table.rows[0].cells):
                    table_shape = shape.table
                    break
            # Route B: Target by index if no name is provided
            else:
                if current_table_idx == table_index:
                    table_shape = shape.table
                    break
                current_table_idx += 1  # Increment if this wasn't the index we wanted
                
    if not table_shape: 
        print(f"Warning: Table at index {table_index} (or name '{table_name}') not found.")
        return

    # ... [The rest of the function remains exactly the same] ...
    num_table_rows = len(table_shape.rows)
    num_table_cols = len(table_shape.columns)
    current_row = start_row

    if include_headers and current_row < num_table_rows:
        for c_idx, col_name in enumerate(df.columns):
            if start_col + c_idx < num_table_cols:
                write_to_cell_preserving_format(table_shape.cell(current_row, start_col + c_idx), col_name)
        current_row += 1

    for r_idx, row_values in enumerate(df.values):
        if current_row >= num_table_rows: break
        for c_idx, val in enumerate(row_values):
            if start_col + c_idx < num_table_cols:
                write_to_cell_preserving_format(table_shape.cell(current_row, start_col + c_idx), val)
        current_row += 1

def write_hierarchical_cell(cell, main_reason, sub_reasons):
    text_frame = cell.text_frame
    p_main = text_frame.paragraphs[0]
    p_main.text = str(main_reason)
    
    ref_font_name, ref_font_size = None, None
    if p_main.runs:
        p_main.runs[0].font.bold = True
        ref_font_name = p_main.runs[0].font.name
        ref_font_size = p_main.runs[0].font.size

    for i, sub_reason in enumerate(sub_reasons, 1):
        p_sub = text_frame.add_paragraph()
        p_sub.text = f"{i}. {sub_reason}"
        p_sub.level = 1 
        if p_sub.runs:
            sub_font = p_sub.runs[0].font
            if ref_font_name: sub_font.name = ref_font_name
            if ref_font_size: sub_font.size = ref_font_size - Pt(2)
            sub_font.bold = False

def fill_complex_reasons_table(slide, list_of_dicts, start_row=1, table_name=None):
    table_shape = None
    for shape in slide.shapes:
        if shape.has_table:
            if table_name:
                if table_name in " ".join(cell.text_frame.text.strip() for cell in shape.table.rows[0].cells):
                    table_shape = shape.table
                    break
            else:
                table_shape = shape.table
                break
                
    if not table_shape: return

    num_table_rows = len(table_shape.rows)
    current_row = start_row

    for row_data in list_of_dicts:
        if current_row >= num_table_rows: break
            
        write_hierarchical_cell(table_shape.cell(current_row, 0), row_data["reason"], row_data.get("sub_reasons", []))
        write_to_cell_preserving_format(table_shape.cell(current_row, 1), f"{row_data['count']} ({row_data['percent']})")
        write_to_cell_preserving_format(table_shape.cell(current_row, 2), str(row_data["delta"]))
        
        current_row += 1

# --- Builder Pipelines ---

def build_single_presentation(template_path, output_path, text_replace_dict, table_data_configs=None, complex_table_configs=None):
    prs = Presentation(template_path)

    # 1. Global Text
    for slide in prs.slides:
        for shape in slide.shapes:
            process_shape(shape, text_replace_dict)

    # 2. Standard DataFrames
    if table_data_configs:
        for config in table_data_configs:
            target_slide = get_slide_by_title(prs, config.get("slide_title"))
            if target_slide:
                fill_table_from_df(
                    slide=target_slide,
                    df=config.get("dataframe"),
                    include_headers=config.get("include_headers", True),
                    start_row=config.get("start_row", 0),
                    start_col=config.get("start_col", 0),
                    table_name=config.get("table_name"),
                    table_index=config.get("table_index", 0)
                )

    # 3. Complex Hierarchical Tables
    if complex_table_configs:
        for config in complex_table_configs:
            target_slide = get_slide_by_title(prs, config.get("slide_title"))
            if target_slide:
                fill_complex_reasons_table(
                    slide=target_slide,
                    list_of_dicts=config.get("complex_data"),
                    start_row=config.get("start_row", 1),
                    table_name=config.get("table_name")
                )

    prs.save(output_path)
    print(f"Success: Saved {output_path}")

def run_batch_presentations(template_path, trials):
    print(f"\nStarting PowerPoint generation for {len(trials)} presentations...")
    for idx, trial in enumerate(trials, 1):
        print(f"Generating PPTX {idx}/{len(trials)}: {trial['output_path']}")
        build_single_presentation(
            template_path=template_path,
            output_path=trial["output_path"],
            text_replace_dict=trial.get("text_replace_dict", {}),
            table_data_configs=trial.get("table_data_configs", []),
            complex_table_configs=trial.get("complex_table_configs", [])
        )
    print("Batch generation complete!")
